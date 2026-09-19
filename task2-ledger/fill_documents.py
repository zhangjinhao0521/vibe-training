"""生成第 2 次课两份 Word 文档的已完成副本。"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "第二次课实验2"
OUTPUT_DIR = SOURCE_DIR / "已完成"
MATERIAL_SOURCE = SOURCE_DIR / "第2次课-实践素材.docx"
OPERATION_SOURCE = SOURCE_DIR / "第2次课-操作文档.docx"
MATERIAL_OUTPUT = OUTPUT_DIR / "第2次课-实践素材-最终版.docx"
OPERATION_OUTPUT = OUTPUT_DIR / "第2次课-操作文档-最终版.docx"

VISION = (
    "我一个人随手在电脑上记录我今天的日常穿搭，月底对我的个人喜好进行总结，"
    "并为下一个月的穿搭做出推荐。"
)


def clear_and_write_cell(cell, text: str, *, size: float = 9.5, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def set_table_repeat_header(table) -> None:
    properties = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def remove_element(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def insert_paragraph_before_table(table, text: str, style=None):
    paragraph = OxmlElement("w:p")
    table._tbl.addprevious(paragraph)
    from docx.text.paragraph import Paragraph

    wrapper = Paragraph(paragraph, table._parent)
    if style is not None:
        wrapper.style = style
    run = wrapper.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return wrapper


def fill_material_document() -> None:
    document = Document(MATERIAL_SOURCE)

    b_prompt = (
        "角色：你是一位资深 Python 工程师。\n"
        "任务：实现一个命令行待办事项应用。\n"
        "上下文：Windows 终端、Python 3.12，单文件 todo.py，供我日常使用。\n"
        "约束：不用第三方库；数据保存到本地 todos.json；界面提示语用中文。\n"
        "示例：输入 add 买牛奶，列表出现“1. [ ] 买牛奶”；输入 done 1 后变成“1. [x] 买牛奶”。\n"
        "验收标准：支持 add/list/done/delete；重启后数据仍在；未知命令给出用法且不崩溃。"
    )
    clear_and_write_cell(document.tables[2].cell(1, 2), b_prompt, size=8.5)

    conclusion = document.paragraphs[8]
    conclusion.text = "我的结论（一句话）：使用六件套提示词后，目标更明确、功能更丰富，AI 的联想内容也更多。"

    document.paragraphs[9].text = "三、需求拆解工作单（穿搭记账本）"
    document.paragraphs[11].text = VISION
    document.paragraphs[18].text = "四、穿搭记账本参考提示词"
    document.paragraphs[21].text = "4.2 穿搭记账本 M1 起步提示词"

    features = [
        ("记录每日穿搭，包括日期、上装、下装、鞋履、颜色、风格、场合、满意度和备注", "M1"),
        ("将记录保存到本地 data.json，程序启动时自动读取", "M1"),
        ("查看全部穿搭记录，并按日期倒序清晰展示", "M2"),
        ("按月份、风格、颜色或场合筛选穿搭记录", "M2"),
        ("统计月度偏好，并根据高满意度记录生成下月穿搭推荐", "M3"),
        ("非法输入兜底；二次确认后删除；空数据友好提示", "M4"),
    ]
    feature_table = document.tables[3]
    for index, (feature, milestone) in enumerate(features, start=1):
        clear_and_write_cell(feature_table.cell(index, 1), feature)
        clear_and_write_cell(feature_table.cell(index, 2), milestone)

    milestones = [
        (
            "最小闭环：命令行记录一套穿搭（日期、单品、颜色、风格、场合、满意度和备注），"
            "保存到本地 JSON，重启后自动读取",
            "① 运行后出现记录/退出菜单 ② 录入后提示已保存，JSON 字段完整 "
            "③ 日期回车默认当天，重启后记录仍在 ④ 日期错误或满意度非 1~5 整数时提示重输",
        ),
        (
            "查看与筛选：列出全部穿搭记录；支持按月份、风格、颜色和场合筛选",
            "① 显示序号、日期、各类单品、颜色、风格、场合、满意度和备注 "
            "② 指定月份只显示该月 ③ 指定风格/颜色/场合只显示匹配记录 "
            "④ 空记录或无匹配结果时友好提示",
        ),
        (
            "月度偏好总结与下月推荐：统计次数、平均满意度、常穿单品/颜色/风格/场合，"
            "并根据高满意度记录生成下一月建议",
            "① 记录数与平均满意度和原始记录一致 ② 各项频次统计正确 "
            "③ 输出下一年月的 1~3 套推荐，单品来自已有记录并说明依据 "
            "④ 无数据月份友好提示、不崩溃",
        ),
        (
            "健壮性与数据管理：非法输入全兜底；按编号删除需确认；损坏数据先备份；空数据提示",
            "① 菜单、日期、评分、月份和删除编号乱输入均不崩溃 "
            "② 删除前询问确认，输入 n 后保留 ③ 输入 y 后删除，重启后结果仍在 "
            "④ 空数据和损坏 JSON 均有友好提示",
        ),
    ]
    milestone_table = document.tables[4]
    for index, (content, acceptance) in enumerate(milestones, start=1):
        clear_and_write_cell(milestone_table.cell(index, 1), content, size=8.3)
        clear_and_write_cell(milestone_table.cell(index, 2), acceptance, size=8.3)
        clear_and_write_cell(milestone_table.cell(index, 3), "☑ 已通过", size=8.3, bold=True)
        prevent_row_split(milestone_table.rows[index])
    set_table_repeat_header(milestone_table)

    # 参考提示词同步到实际完成的穿搭项目，方便提交最满意的一条提示词。
    prompts = [
        (
            "你是一位资深 Python 工程师。\n"
            "任务：实现个人穿搭记账本 M1——最小闭环。\n"
            "上下文：Windows 终端、Python 3.12，单文件 ledger.py，数据存 data.json。\n"
            "功能：菜单为“1 记录今日穿搭 / 2 退出”；依次输入日期、上装、下装、鞋履、"
            "颜色、风格、场合、满意度和备注并保存；启动时自动读取已有数据。\n"
            "约束：只用标准库；UTF-8 JSON；中文提示；日期回车默认今天；满意度限 1~5。\n"
            "验收：重启后数据仍在；错误日期和评分提示重输；必要字段空回车不崩溃。"
        ),
        (
            "在现有 ledger.py 基础上，只增加“查看与筛选穿搭”功能，不改动新增和存储逻辑：\n"
            "菜单增加查看项；全部记录按日期倒序显示完整字段；支持按月份、风格、颜色和场合筛选；"
            "空数据或无匹配结果时友好提示。改完列出修改的函数和自测方法。"
        ),
        (
            "在现有 ledger.py 基础上，只增加“月度偏好总结与下月推荐”功能，不改动其他功能：\n"
            "输入年月（回车默认本月），统计记录次数、平均满意度、常穿风格/颜色/上装/下装/鞋履/场合；"
            "从本月真实记录中按满意度和使用次数生成下一月 1~3 套推荐，并说明依据；无数据月份友好提示。"
        ),
        (
            "请不要修改任何代码，只做两件事：\n"
            "1. 用一段话解释 ledger.py 的结构和数据流向，从输入到写入 data.json。\n"
            "2. 逐行解释 load_records() 和 save_records() 的核心逻辑。\n"
            "解释后出 2 道小问题考我，并在我回答后判分纠错。"
        ),
    ]
    for table_index, prompt in zip((6, 7, 8, 9), prompts):
        clear_and_write_cell(document.tables[table_index].cell(0, 0), prompt, size=9)

    template_table = document.tables[11]
    template_element = template_table._tbl
    body = template_element.getparent()
    template_position = body.index(template_element)

    acceptance_sections = [
        (
            "M1 验收清单",
            [
                ("正常记录并持久化", "新增一套后退出重启，检查 JSON 字段和记录仍完整"),
                ("日期与满意度校验", "日期输错；满意度输入字母、0、6，确认提示重输"),
                ("边界情况：满意度 1/5、超长备注", "分别录入一次并重启查看"),
                ("异常输入：必要字段空回车", "实际操作一次，确认提示重输且不崩溃"),
                ("改动范围受控", "核对本阶段只有新增、读取、保存和退出"),
            ],
        ),
        (
            "M2 验收清单",
            [
                ("全部记录显示完整", "预置跨月记录，检查日期、单品、偏好字段与排序"),
                ("筛选结果准确", "分别按月份、风格、颜色和场合筛选并人工核对"),
                ("边界情况：无匹配结果", "筛选不存在的月份和关键词"),
                ("异常输入：筛选方式乱输入", "输入字母、空回车、越界数字，确认不崩溃"),
                ("改动范围受控：M1 回归", "重新新增记录并重启，确认存储功能正常"),
            ],
        ),
        (
            "M3 验收清单",
            [
                ("月度汇总准确", "用 3 条 9 月数据人工核对数量、平均分和偏好频次"),
                ("下月推荐可追溯", "确认输出 1~3 套，单品来自 9 月并显示推荐依据"),
                ("边界情况：12 月跨年", "汇总 2026-12，确认推荐月份为 2027-01"),
                ("异常输入：月份错误/无数据", "输入 2026-13 和空数据月份，确认友好提示"),
                ("改动范围受控：M1~M2 回归", "重新测试新增、读取、完整查看与筛选"),
            ],
        ),
        (
            "M4 验收清单",
            [
                ("删除需二次确认", "同一记录先输入 n 验证保留，再输入 y 验证删除"),
                ("删除结果持久化", "确认删除后重启，验证记录不再出现"),
                ("边界情况：空数据", "在空数据下执行查看、汇总和删除"),
                ("异常输入与损坏 JSON", "菜单/编号乱输入；使用损坏 JSON 启动并检查备份"),
                ("改动范围受控：M1~M3 回归", "运行全部 6 项自动测试和完整交互测试"),
            ],
        ),
    ]

    # 删除原来的通用空表，换成每个里程碑一张已填写清单。
    remove_element(template_element)
    insertion_index = template_position
    caption_style = document.paragraphs[32].style
    for title, rows in acceptance_sections:
        paragraph = OxmlElement("w:p")
        body.insert(insertion_index, paragraph)
        from docx.text.paragraph import Paragraph

        caption = Paragraph(paragraph, document._body)
        caption.style = caption_style
        run = caption.add_run(title)
        run.bold = True
        insertion_index += 1

        table_copy = deepcopy(template_table._tbl)
        body.insert(insertion_index, table_copy)
        insertion_index += 1
        from docx.table import Table

        table = Table(table_copy, document._body)
        for row_index, (item, method) in enumerate(rows, start=1):
            clear_and_write_cell(table.cell(row_index, 0), item, size=8.5)
            clear_and_write_cell(table.cell(row_index, 1), method, size=8.5)
            clear_and_write_cell(table.cell(row_index, 2), "☑", size=10, bold=True)
            prevent_row_split(table.rows[row_index])
        set_table_repeat_header(table)

        spacer = OxmlElement("w:p")
        body.insert(insertion_index, spacer)
        insertion_index += 1

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(MATERIAL_OUTPUT)


def fill_operation_document() -> None:
    document = Document(OPERATION_SOURCE)

    document.paragraphs[16].text = "三、实验 2：个人穿搭记账本（80 分钟）"
    document.paragraphs[18].text = f"写一句话愿景：{VISION}"

    for index in (7, 8, 9, 10):
        document.paragraphs[index].text = document.paragraphs[index].text.replace("□", "☑", 1)

    document.paragraphs[26].text = (
        "验收时至少覆盖三类情况：正常路径（完整记录并重启查看）、边界情况"
        "（满意度 1/5、超长备注、12 月跨年）和异常输入（错误日期、字母、空回车、越界值）。"
        "以上均已通过。"
    )

    # 这三项已经在当前工作区完成或生成；作品墙属于课外发布行为，不代用户勾选。
    document.paragraphs[38].text = (
        "☑  穿搭记账本代码（task2-ledger 文件夹）+ git log --oneline 记录（逐里程碑提交）。"
    )
    document.paragraphs[39].text = "☑  《提示词对比实验记录单》电子版已填写。"
    document.paragraphs[40].text = (
        "□  你认为写得最得意的一条提示词，贴到班级作品墙（待本人发布）。"
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(OPERATION_OUTPUT)


if __name__ == "__main__":
    fill_material_document()
    fill_operation_document()
    print(MATERIAL_OUTPUT)
    print(OPERATION_OUTPUT)
