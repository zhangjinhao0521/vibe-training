from __future__ import annotations

from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SHOT_DIR = ROOT / "截图"
DOC_SHOT_DIR = SHOT_DIR / "文档用"
OUTPUT = ROOT / "第2次课实验最终总结.docx"

NAVY = "142331"
INK = "17212B"
TEAL = "0B7A75"
GREEN = "16A36A"
LIME = "D9FF35"
LIGHT = "F2F5F7"
MID = "D7DEE4"
MUTED = "66717D"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID, size=5) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_run_font(run, name="Microsoft YaHei", size=None, color=None, bold=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_text(paragraph, text, *, bold=False, color=INK, size=10.5, font="Microsoft YaHei"):
    run = paragraph.add_run(text)
    set_run_font(run, font, size, color, bold)
    return run


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for style_name, size, color in (
        ("Title", 30, NAVY),
        ("Heading 1", 20, NAVY),
        ("Heading 2", 14, TEAL),
        ("Heading 3", 11.5, NAVY),
    ):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    caption = document.styles["Caption"]
    caption.font.name = "Microsoft YaHei"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    caption.font.size = Pt(8.5)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.font.italic = False
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(7)
    caption.paragraph_format.keep_with_next = True

    for current_section in document.sections:
        header = current_section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        add_text(paragraph, "AI 应用开发实训  ·  第 2 次课  ·  实验最终总结", color=MUTED, size=8.5)

        footer = current_section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_after = Pt(0)
        add_text(fp, "—  ", color=MUTED, size=8)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        fp._p.append(field)
        add_text(fp, "  —", color=MUTED, size=8)

    props = document.core_properties
    props.title = "第2次课实验最终总结——从结构化提示词到个人穿搭记账本"
    props.subject = "AI 应用开发实训：实验 1 与实验 2 总结"
    props.keywords = "Vibe Coding, 六件套提示词, 穿搭记账本, Python, Git"


def add_page_break(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_section_heading(document: Document, number: str, title: str, subtitle: str = "") -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(1.35)
    table.columns[1].width = Cm(15.5)
    table.cell(0, 0).width = Cm(1.35)
    table.cell(0, 1).width = Cm(15.5)
    set_table_borders(table, WHITE, 0)
    num_cell, title_cell = table.rows[0].cells
    set_cell_shading(num_cell, LIME)
    set_cell_margins(num_cell, 110, 80, 100, 80)
    set_cell_margins(title_cell, 0, 220, 0, 0)
    num_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    title_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = num_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    add_text(p, number, bold=True, color=NAVY, size=15)
    p = title_cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_text(p, title, bold=True, color=NAVY, size=20)
    if subtitle:
        p = title_cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        add_text(p, subtitle, color=MUTED, size=8.5)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(document: Document, text: str, label: str | None = None, fill=LIGHT, accent=TEAL) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(16.8)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 160, 220, 160, 220)
    set_table_borders(table, accent, 8)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if label:
        add_text(p, label + "  ", bold=True, color=accent, size=9.5)
    add_text(p, text, color=INK, size=10.5)


def add_body(document: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_prefix and text.startswith(bold_prefix):
        add_text(p, bold_prefix, bold=True)
        add_text(p, text[len(bold_prefix):])
    else:
        add_text(p, text)


def add_bullets(document: Document, items: list[str], *, compact=False) -> None:
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(2 if compact else 4)
        p.paragraph_format.line_spacing = 1.25
        add_text(p, item, size=9.7 if compact else 10.2)


def add_caption(document: Document, text: str) -> None:
    p = document.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    p.add_run(text)


def add_image(document: Document, path: Path, width: float, caption: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width))
    add_caption(document, caption)


def set_cell_text(cell, text: str, *, bold=False, color=INK, size=9, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    add_text(p, text, bold=bold, color=color, size=size)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None, font_size=8.8):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        set_cell_shading(cell, NAVY)
        set_cell_text(cell, header, bold=True, color=WHITE, size=8.8, align=WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            cell.width = Cm(widths[index])
    for row_index, row_data in enumerate(rows, start=1):
        cells = table.add_row().cells
        prevent_row_split(table.rows[row_index])
        for index, value in enumerate(row_data):
            if row_index % 2 == 0:
                set_cell_shading(cells[index], "F7F9FA")
            set_cell_text(cells[index], value, size=font_size)
            if widths:
                cells[index].width = Cm(widths[index])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def prepare_images() -> dict[str, Path]:
    DOC_SHOT_DIR.mkdir(parents=True, exist_ok=True)
    crop_bottom = {
        "实验2-M1M2记录与筛选.png": 690,
        "实验2-M3月度总结与推荐.png": 755,
        "实验2-M4删除确认.png": 755,
        "实验2-自动测试通过.png": 555,
        "Git-里程碑提交记录.png": 470,
    }
    results = {}
    for source in SHOT_DIR.glob("*.png"):
        target = DOC_SHOT_DIR / source.name
        image = Image.open(source).convert("RGB")
        if source.name in crop_bottom:
            image = image.crop((0, 0, image.width, min(crop_bottom[source.name], image.height)))
        image.save(target, quality=95)
        results[source.name] = target
    return results


def build_report() -> Path:
    images = prepare_images()
    document = Document()
    configure_document(document)

    # Cover
    document.add_paragraph().paragraph_format.space_after = Pt(24)
    tag = document.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(tag, "V I B E   C O D I N G   T R A I N I N G", bold=True, color=TEAL, size=10)
    document.add_paragraph().paragraph_format.space_after = Pt(18)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    add_text(title, "第 2 次课实验最终总结", bold=True, color=NAVY, size=31)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    add_text(subtitle, "从结构化提示词到个人穿搭记账本", bold=True, color=TEAL, size=17)

    band = document.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    band.autofit = False
    band.columns[0].width = Cm(13.5)
    set_table_borders(band, LIME, 0)
    cell = band.cell(0, 0)
    set_cell_shading(cell, LIME)
    set_cell_margins(cell, 180, 260, 180, 260)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "提示词对比实验  +  穿搭记账本 M1～M4", bold=True, color=NAVY, size=12)

    document.add_paragraph().paragraph_format.space_after = Pt(64)
    meta = add_table(
        document,
        ["课程", "模块", "实验日期"],
        [["AI 应用开发实训", "模块一 · 第 2 次课", "2026 年 9 月 11 日"]],
        [5.6, 5.6, 5.6],
        9.3,
    )
    for cell in meta.rows[1].cells:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_before = Pt(22)
    add_text(info, "姓名：________________    学号：________________    班级：________________", color=MUTED, size=10)
    document.add_paragraph().paragraph_format.space_after = Pt(40)
    quote = document.add_paragraph()
    quote.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(quote, "“先把目标说清楚，再把结果一步步做实。”", bold=True, color=NAVY, size=12)
    add_page_break(document)

    # Page 2
    add_section_heading(document, "01", "课程目标与产出总览", "依据两份最终版课程文档与当前可核验产物整理")
    add_body(document, "本次实操共 120 分钟。实验 1 用提示词对比验证“说清楚”的价值；实验 2 完整走过需求拆解、分步实现、行为验收和 Git 留痕。最终产物既包含可运行程序，也包含工作单、测试证据和本总结报告。")
    add_table(
        document,
        ["环节", "时间", "核心问题", "可核验产物", "完成情况"],
        [
            ["实验 1", "40 min", "结构化提示词是否让结果更贴近目标？", "《街头霸王 6》安利网页、桌面与手机截图、A/B 自评记录", "网页正常渲染；自评 2.0 → 4.5"],
            ["实验 2", "80 min", "如何把愿景拆成可验收的垂直切片？", "Python 穿搭记账本、JSON 数据、6 项自动测试、M1～M4 清单", "M1～M4 功能均通过验收"],
        ],
        [2.1, 1.6, 4.5, 5.4, 3.2],
        8.3,
    )
    document.add_heading("统一工作方法", level=2)
    workflow = document.add_table(rows=1, cols=5)
    workflow.alignment = WD_TABLE_ALIGNMENT.CENTER
    workflow.autofit = False
    set_table_borders(workflow, WHITE, 0)
    for index, text in enumerate(("明确任务", "补足上下文", "写验收标准", "逐里程碑实现", "测试并提交")):
        cell = workflow.cell(0, index)
        set_cell_shading(cell, LIME if index in (0, 2, 4) else LIGHT)
        set_cell_text(cell, f"{index + 1:02d}\n{text}", bold=True, color=NAVY, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell.width = Cm(3.25)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    document.add_heading("最终交付清单", level=2)
    add_bullets(document, [
        "实验 1：响应式网页 index.html、样式与交互脚本、评测原文及 4 张视觉素材。",
        "实验 2：ledger.py、data.json、README、验收记录和 test_ledger.py。",
        "课程文档：操作文档最终版、实践素材最终版；工作单中的 M1～M4 均已填写并勾选。",
        "过程证据：7 张对应截图，覆盖页面、运行输出、自动测试和 Git 记录。",
    ], compact=True)
    add_callout(document, "本文只对现有文件、实际运行输出和 Git 历史作结论；演示数据、自评数据与正式个人数据分别标注。", "说明")
    add_page_break(document)

    # Page 3
    add_section_heading(document, "02", "实验 1 · 六件套提示词成果", "当前可核验作品：《街头霸王 6》论坛安利 / 评测网页")
    add_body(document, "六件套提示词从角色、任务、上下文、约束、示例和验收标准六个维度约束输出。最终网页以硬核玩家为受众，采用论坛式口吻组织“省流总结、核心亮点、劝退点、入坑指南”，并保留真实优缺点分析。")
    add_image(document, images["实验1-街霸6安利页面.png"], 16.7, "图 1  实验 1 桌面端首屏：品牌、主题、文章结构与真实游戏素材完整呈现")
    add_callout(document, "记录单中的 A/B 两组均为 1 轮；可用度自评由 2.0/5 提升至 4.5/5，提高 2.5 分。该分数属于实验自评，不等同于外部测评。", "记录结果")
    add_page_break(document)

    # Page 4
    add_section_heading(document, "03", "实验 1 · 质量验收", "重点检查内容结构、游戏语境与移动端可读性")
    layout = document.add_table(rows=1, cols=2)
    layout.alignment = WD_TABLE_ALIGNMENT.CENTER
    layout.autofit = False
    set_table_borders(layout, WHITE, 0)
    left, right = layout.rows[0].cells
    left.width = Cm(7.1)
    right.width = Cm(9.5)
    set_cell_margins(left, 0, 0, 0, 150)
    set_cell_margins(right, 0, 250, 0, 0)
    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.add_run().add_picture(str(images["实验1-手机端页面.png"]), width=Cm(6.5))
    rp = right.paragraphs[0]
    add_text(rp, "验收观察", bold=True, color=NAVY, size=14)
    observations = [
        "结构完整：首屏明确给出主题和推荐结论，正文覆盖亮点、缺点与入坑路径。",
        "语言贴合受众：包含打击感、立回、确反、回滚网络码、资源管理等玩家术语。",
        "移动适配：390 × 844 视口无横向溢出，标题、导航和评分信息均可读。",
        "交互完整：提供深浅主题、摘要复制、读者路线切换、阅读进度和返回顶部。",
        "视觉素材真实：4 张图片均正常加载，内容与游戏主体一致。",
    ]
    for item in observations:
        p = right.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.first_line_indent = Cm(-0.28)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.25
        add_text(p, "✓ ", bold=True, color=GREEN, size=10)
        add_text(p, item, size=9.8)
    add_caption(document, "图 2  390 × 844 手机端首屏；自动检查结果：无横向溢出")
    add_callout(document, "六件套的核心价值不是让页面“更长”，而是让内容形态、语气、必须包含的部分和验收方法在生成前就被共同约定。", "实验结论")
    add_body(document, "需要如实说明：操作文档设计的是两组“待办事项应用”对比，工作单保留了 A/B 自评数据；当前仓库中可直接核验的实验 1 成品则是自定义六件套任务生成的《街头霸王 6》网页，未发现两组待办成品截图。")
    add_page_break(document)

    # Page 5
    add_section_heading(document, "04", "实验 2 · 愿景与需求拆解", "把“穿搭记录与推荐”拆成四个可独立验收的垂直切片")
    add_callout(document, "我一个人随手在电脑上记录我今天的日常穿搭，月底对我的个人喜好进行总结，并为下一个月的穿搭做出推荐。", "一句话愿景", fill="ECF8F3", accent=GREEN)
    add_table(
        document,
        ["里程碑", "可用增量", "关键验收点", "状态"],
        [
            ["M1", "记录日期、单品、颜色、风格、场合、满意度和备注；保存到本地 JSON", "错误日期/评分重输；重启后数据仍在", "已通过"],
            ["M2", "查看全部记录；按月份、风格、颜色或场合筛选", "倒序展示字段完整；空数据与无匹配有提示", "已通过"],
            ["M3", "月度次数、平均满意度和偏好统计；生成下月 1～3 套建议", "结果可人工核对；无数据月份不崩溃", "已通过"],
            ["M4", "输入兜底、删除二次确认、损坏 JSON 备份", "取消删除保留；确认删除持久化；异常不崩溃", "已通过"],
        ],
        [1.4, 6.0, 7.0, 2.2],
        8.4,
    )
    document.add_heading("文件结构与数据流", level=2)
    add_callout(document, "终端输入  →  输入校验  →  内存中的记录列表  →  UTF-8 JSON 临时文件  →  原子替换 data.json", "数据流")
    add_bullets(document, [
        "应用只使用 Python 标准库，数据完全保存在本机，不依赖网络服务。",
        "保存时执行 flush / fsync，并用临时文件替换正式文件，降低中断造成的数据损坏风险。",
        "加载时逐条校验字段；JSON 损坏时先复制为 data.invalid-时间.json，再从空记录启动。",
        "推荐采用可解释规则：按完整搭配的平均满意度、出现次数和稳定顺序排序，选取前 1～3 套。",
    ], compact=True)
    add_page_break(document)

    # Page 6
    add_section_heading(document, "05", "实验 2 · M1 与 M2", "完成“记录—保存—重启读取—查看筛选”的最小闭环")
    add_image(document, images["实验2-M1M2记录与筛选.png"], 16.7, "图 3  三条验收用穿搭记录从 JSON 读取，并按 2026-09 筛选后倒序展示")
    add_table(
        document,
        ["验收项", "实际结果", "结论"],
        [
            ["记录完整性", "日期、上装、下装、鞋履、配色、风格、场合、满意度、备注均写入", "通过"],
            ["持久化", "保存后退出并重新运行，记录仍可读取", "通过"],
            ["筛选与排序", "支持月份/风格/颜色/场合；结果按日期倒序", "通过"],
            ["异常与空数据", "必要字段为空、错误选项和无匹配结果均给出中文提示", "通过"],
        ],
        [4.0, 10.2, 2.4],
        8.6,
    )
    add_callout(document, "图中三条记录是为验收构造的演示数据，保存在截图目录的 demo-data.json；正式 data.json 当前为空数组，不代表已经记录了真实整月生活数据。", "数据口径")
    add_page_break(document)

    # Page 7
    add_section_heading(document, "06", "实验 2 · M3 月度总结与推荐", "从可核对的历史记录中提取偏好，并给出下一月建议")
    add_image(document, images["实验2-M3月度总结与推荐.png"], 16.7, "图 4  2026 年 9 月偏好汇总与 2026 年 10 月规则推荐")
    add_table(
        document,
        ["统计维度", "验收结果"],
        [
            ["记录次数 / 平均满意度", "3 次 / 4.0 分"],
            ["风格 / 颜色", "休闲 2 次、通勤 1 次 / 白色 2 次、蓝色 1 次"],
            ["高频单品", "白T 2 次、牛仔裤 2 次、小白鞋 2 次"],
            ["下一月推荐", "按满意度 5.0、4.0、3.0 排出 3 套，并逐条说明依据"],
        ],
        [5.0, 11.6],
        8.8,
    )
    add_callout(document, "该功能是基于本月已穿完整搭配的可解释规则推荐：不会凭空生成新单品，也不是机器学习预测。数据越完整，月末总结越有参考价值。", "算法边界")
    add_page_break(document)

    # Page 8
    add_section_heading(document, "07", "实验 2 · M4 与质量验证", "删除需确认；自动测试覆盖关键数据路径")
    add_image(document, images["实验2-M4删除确认.png"], 16.3, "图 5  同一记录先取消、再确认删除；最终数据文件复查为空数组")
    add_image(document, images["实验2-自动测试通过.png"], 16.3, "图 6  6 项自动测试全部通过（实际运行输出统一排版）")
    add_body(document, "自动测试覆盖保存/读取、筛选与排序、月度汇总和推荐、空月份、取消/确认删除及损坏 JSON 备份；人工交互验收另覆盖日期错误、评分边界 1/5、超长备注、菜单乱输入和 2026-12 → 2027-01 跨年。")
    add_page_break(document)

    # Page 9
    add_section_heading(document, "08", "Git 留痕、复盘与结论", "以提交记录证明增量开发，以验收结果约束最终结论")
    add_image(document, images["Git-里程碑提交记录.png"], 16.5, "图 7  Git 记录：M1、M2、M3 独立提交，M4 与完整材料进入最终整合提交")
    add_table(
        document,
        ["提交", "内容", "说明"],
        [
            ["be23d92", "M1：最小闭环可记录穿搭", "独立里程碑"],
            ["84bad64", "M2：查看和筛选穿搭记录", "独立里程碑"],
            ["a1ec523", "M3：月度偏好总结与推荐", "独立里程碑"],
            ["e5b7667", "最终整合：M4、测试、说明、实验 1 与已填文档", "提交信息为“9.11”"],
        ],
        [3.0, 9.3, 4.3],
        8.5,
    )
    document.add_heading("复盘", level=2)
    add_bullets(document, [
        "结构化提示词先定义目标、边界和验收标准，显著减少“做出来但不是想要的”这种偏差。",
        "垂直切片让每一步都有完整可用结果；回归测试保证增加新功能时不破坏旧功能。",
        "本地 JSON、异常备份和删除确认把“能运行”提高到“可放心使用”；Git 让过程可追溯。",
        "仍可改进：以后应保存实验 1 两组原始提示词与原始截图；M4 应形成独立命名提交；真实使用后再评估偏好推荐。",
    ], compact=True)
    add_callout(document, "本次实验最终完成了从“把需求说清楚”到“把产品做完整、测清楚、留证据”的闭环。穿搭记账本达到课程 M4 优秀档功能要求；所有结论均能在代码、测试、截图或 Git 历史中找到对应依据。", "最终结论", fill="ECF8F3", accent=GREEN)

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
